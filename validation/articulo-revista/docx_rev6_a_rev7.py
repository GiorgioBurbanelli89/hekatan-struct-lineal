# rev6 -> rev7 del artículo de la revista
import copy, re
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

SRC = r"C:/Users/j-b-j/Downloads/Articulo_Revista_Politecnica_Hekatan_rev6_SAP.docx"
DST = r"C:/Users/j-b-j/Downloads/Articulo_Revista_Politecnica_Hekatan_rev7.docx"
XSL = etree.XSLT(etree.parse(r"C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL"))

d = Document(SRC)
log = []

def P():
    return d.paragraphs

def find(start):
    hits = [p for p in P() if p.text.strip().startswith(start)]
    assert len(hits) == 1, (start, len(hits))
    return hits[0]

def rep(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1); log.append(f"OK  {old[:40]!r}"); return
    raise SystemExit(f"NO ENCONTRADO en un run: {old!r}")

def delete(p):
    p._element.getparent().remove(p._element)

def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)

# ---- 3. Burbano 2024 -> 2026
rep(find("Hekatan Struct se diseñó"), "Burbano, 2024", "Burbano, 2026")

# ---- 4. citar MIDUVI 2015c y CSI 2017
rep(find("El análisis de cimentaciones superficiales"),
    "El análisis de cimentaciones superficiales se aborda",
    "El análisis de cimentaciones superficiales, regulado por la NEC-SE-GC (MIDUVI, 2015c), se aborda")
rep(find("El contraste con ETABS exige"),
    "documentadas por él mismo.",
    "documentadas por él mismo (Computers and Structures, Inc., 2017).")

# ---- 2. Bowles en Referencias (tras Batoz)
batoz = find("Batoz, J. L.")
bw = copy.deepcopy(batoz._element)
batoz._element.addnext(bw)
from docx.text.paragraph import Paragraph
bwp = Paragraph(bw, batoz._parent)
rs = bwp.runs
rs[0].text = "Bowles, J. E. (1996). "
rs[1].text = "Foundation analysis and design"
rs[2].text = " (5.ª ed.). McGraw-Hill."
for r in rs[3:]:
    r._element.getparent().remove(r._element)
log.append("OK  Bowles añadido")

# ---- 9. biografía de Dámaso
bios = [p for p in P() if p.text.startswith("Zambrano Mendoza, Dámaso. Ingeniero")]
assert len(bios) == 2
rep(bios[0], " [REVISAR Y COMPLETAR CON SUS DATOS REALES]", "")
delete(bios[1]); log.append("OK  bio duplicada borrada")

# ---- 16. gramática
rep(find("Los resultados evidencian que una plataforma"),
    "y no compromete la validez", "y no comprometen la validez")

# ---- 11. introducción: 3 párrafos -> 1
a = find("Frente a esta problemática")
b = find("En el ámbito del código abierto existen")
c = find("A escala internacional, la migración")
set_text(a,
 "Frente a esta problemática, las tecnologías de código abierto y la ejecución de código nativo en el "
 "navegador mediante WebAssembly abren la posibilidad de democratizar el cálculo estructural sin sacrificar "
 "precisión ni rendimiento (Haas et al., 2017). En el ámbito del código abierto existen solvers de elementos "
 "finitos consolidados —como OpenSees, orientado a la ingeniería sísmica, Code_Aster o Frame3DD—, además de "
 "numerosas bibliotecas de álgebra lineal y mallado; sin embargo, la mayoría opera como aplicaciones de "
 "escritorio o de línea de comandos que exigen instalación, compilación o conocimientos de programación, y rara "
 "vez integran de forma nativa los parámetros de una normativa local como la NEC-15. Aunque WebAssembly, junto "
 "con motores de visualización como WebGL, ha habilitado desde entornos de cálculo científico hasta simuladores "
 "interactivos con un rendimiento cercano al nativo, su aplicación al análisis estructural conforme a normativas "
 "locales sigue siendo incipiente, oportunidad tecnológica que aprovecha el presente trabajo.")
delete(b); delete(c); log.append("OK  introducción fusionada")

# ---- 1. cortante en Y
p = find("La carga sísmica reactiva y el cortante estático coincidieron")
rep(p, "En los dos programas la relación Vdin/Vest supera el 80 % que la NEC-15-SE-DS exige para estructuras regulares (§6.2.2.b).",
 "En la dirección X la relación Vdin/Vest es del 82.6 % (82.7 % en SAP2000) y supera el 80 % que la NEC-15-SE-DS "
 "exige para estructuras regulares (§6.2.2.b). En la dirección Y, en cambio, es del 51.0 % (50.9 % en SAP2000). "
 "La causa no es una discrepancia entre programas, pues ambos obtienen el mismo valor, sino una propiedad de la "
 "estructura: en Y la masa se reparte entre dos modos acoplados con torsión —el segundo, con el 43.0 % de la masa, "
 "y el quinto, con el 38.7 %, según SAP2000—; el quinto (T = 0.118 s) cae en la rama ascendente del espectro, por "
 "debajo de T0 ≈ 0.30 s, donde Sa ≈ 0.53 en lugar de 0.72, y la combinación CQC compone las respuestas modales en "
 "forma cuadrática en lugar de sumarlas. Conforme a la NEC-15-SE-DS (§6.2.2.b), el cortante dinámico en Y debe "
 "escalarse hasta el 80 % del estático, con un factor 0.80 · 32.07 / 16.37 = 1.57.")

t4 = [t for t in d.tables if t.rows[0].cells[0].text.strip() == "Concepto"][0]
last = t4.rows[-1]._tr
newtr = copy.deepcopy(last); last.addnext(newtr)
from docx.table import _Row
nr = _Row(newtr, t4)
for cell, txt in zip(nr.cells, ["Vdin,Y / Vest", "51.0 %", "50.9 %", "< 80 %: escalar × 1.57"]):
    ps = cell.paragraphs; set_text(ps[0], txt)
log.append("OK  fila Vdin,Y/Vest en Tabla 4")

# ---- 13. suite
rep(find("Estos contrastes se apoyan en un registro"),
    "superó 608 de 617 comprobaciones automáticas,",
    "superó 613 de 622 comprobaciones automáticas —las nueve restantes corresponden a funciones ajenas al caso "
    "aquí validado, como el mallado automático de paños irregulares, el muelle de área frente a SAFE y las fuerzas "
    "de cáscara en los nudos, y se mantienen abiertas sin relajar sus tolerancias—,")

# ---- 14. Tabla 7 en %
t7 = [t for t in d.tables if t.rows[0].cells[0].text.strip() == "Caso de referencia"][0]
set_text(t7.rows[0].cells[2].paragraphs[0], "Diferencia (%)")
set_text(t7.rows[1].cells[2].paragraphs[0], "0.00")
set_text(t7.rows[2].cells[2].paragraphs[0], "−1.0 a +0.3")
set_text(t7.rows[3].cells[2].paragraphs[0], "< 0.33")
set_text(t7.rows[4].cells[2].paragraphs[0], "0.00")
set_text(t7.rows[5].cells[2].paragraphs[0], "< 1")
rep(find("La concordancia obtenida"), "del orden de la unidad en pórticos", "exacta en pórticos")

# ---- 15. mover Ec. (7) tras el párrafo que la presenta
eq7 = [p for p in P() if p.text.rstrip().endswith("(7)")][0]
find("A partir del espectro, el cortante basal")._element.addnext(eq7._element)
log.append("OK  Ec. (7) movida")

# ---- 10. ecuaciones -> OMML
M = 'xmlns="http://www.w3.org/1998/Math/MathML"'
def mi(x): return f"<mi>{x}</mi>"
def mo(x): return f"<mo>{x}</mo>"
def sub(a, b): return f"<msub><mi>{a}</mi><mi>{b}</mi></msub>"
def br(o, inner, c): return f'<mo stretchy="false">{o}</mo>{inner}<mo stretchy="false">{c}</mo>'
K = br("[", mi("K"), "]"); Mm = br("[", mi("M"), "]")
phi_i = sub("φ", "i")
phiT = "<msubsup><mi>φ</mi><mi>i</mi><mi>T</mi></msubsup>"
integ = "<munderover><mo>∫</mo><mrow><mo>−</mo><mi>t</mi><mo>/</mo><mn>2</mn></mrow><mrow><mi>t</mi><mo>/</mo><mn>2</mn></mrow></munderover>"
EQ = {
 1: K + br("{", mi("u"), "}") + mo("=") + br("{", mi("f"), "}"),
 2: mi("φ") + mo("=") + "<mfrac><mrow><mn>12</mn><mi>E</mi><mi>I</mi></mrow><mrow><mi>G</mi>" + sub("A", "s") + "<msup><mi>L</mi><mn>2</mn></msup></mrow></mfrac>",
 3: sub("F", "ij") + mo("=") + integ + sub("σ", "ij") + "<mi>d</mi><mi>z</mi>" + mo(";") + "<mspace width='1em'/>"
    + sub("M", "ij") + mo("=") + integ + sub("σ", "ij") + "<mi>z</mi><mi>d</mi><mi>z</mi>" + mo(";") + "<mspace width='1em'/>"
    + sub("V", "i3") + mo("=") + integ + sub("τ", "i3") + "<mi>d</mi><mi>z</mi>",
 4: br("(", K + mo("−") + "<msubsup><mi>ω</mi><mi>i</mi><mn>2</mn></msubsup>" + Mm, ")") + br("{", phi_i, "}") + mo("=") + br("{", "<mn>0</mn>", "}"),
 5: sub("Γ", "i") + mo("=") + "<mfrac><msup><mrow>" + br("(", phiT + Mm + br("{", mi("r"), "}"), ")") + "</mrow><mn>2</mn></msup><mrow>"
    + br("(", phiT + Mm + br("{", phi_i, "}"), ")") + br("(", "<msup><mrow>" + br("{", mi("r"), "}") + "</mrow><mi>T</mi></msup>" + Mm + br("{", mi("r"), "}"), ")")
    + "</mrow></mfrac>" + mo(";") + "<mspace width='1em'/>" + "<munder><mo>∑</mo><mi>i</mi></munder>" + sub("Γ", "i") + mo("≥") + "<mn>0.90</mn>",
 6: "<msub><mi>S</mi><mi>a</mi></msub>" + mo("=") + "<mrow><mo>{</mo><mtable columnalign='left'>"
    "<mtr><mtd><mi>Z</mi>" + sub("F", "a") + br("[", "<mn>1</mn>" + mo("+") + br("(", mi("η") + mo("−") + "<mn>1</mn>", ")") + "<mfrac><mi>T</mi>" + sub("T", "0") + "</mfrac>", "]") + "</mtd><mtd><mn>0</mn>" + mo("≤") + mi("T") + mo("&lt;") + sub("T", "0") + "</mtd></mtr>"
    "<mtr><mtd><mi>η</mi><mi>Z</mi>" + sub("F", "a") + "</mtd><mtd>" + sub("T", "0") + mo("≤") + mi("T") + mo("≤") + sub("T", "c") + "</mtd></mtr>"
    "<mtr><mtd><mi>η</mi><mi>Z</mi>" + sub("F", "a") + "<msup><mrow>" + br("(", "<mfrac>" + sub("T", "c") + "<mi>T</mi></mfrac>", ")") + "</mrow><mi>r</mi></msup></mtd><mtd>" + mi("T") + mo("&gt;") + sub("T", "c") + "</mtd></mtr>"
    "</mtable></mrow>",
 7: mi("V") + mo("=") + "<mfrac><mrow><mi>I</mi><mo>·</mo><msub><mi>S</mi><mi>a</mi></msub>" + br("(", mi("T"), ")") + "</mrow><mrow><mi>R</mi><mo>·</mo>" + sub("φ", "P") + "<mo>·</mo>" + sub("φ", "E") + "</mrow></mfrac><mo>·</mo><mi>W</mi>",
 8: mi("p") + mo("=") + sub("k", "s") + "<mo>·</mo>" + mi("w"),
}
OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
for n, body in EQ.items():
    ps = [p for p in P() if re.search(rf"\t\({n}\)\s*$", p.text)]
    assert len(ps) == 1, (n, len(ps))
    p = ps[0]
    mml = etree.fromstring(f"<math {M}><mrow>{body}</mrow></math>")
    omml = XSL(mml).getroot()
    # dejar: run "\t" + oMath + run "\t(n)"
    runs = p.runs
    rpr = runs[-1]._element.find(qn("w:rPr"))
    for r in runs:
        r._element.getparent().remove(r._element)
    def run(txt):
        r = p.add_run(txt)
        if rpr is not None:
            r._element.insert(0, copy.deepcopy(rpr))
            r.italic = False
        return r
    run("\t")
    p._element.append(omml)
    run(f"\t({n})")
    log.append(f"OK  Ec. ({n}) -> OMML")

# párrafo que describe la Ec. (6): rama ascendente y T0
rep(find("La plataforma incorpora directamente los parámetros sísmicos"),
    "donde Fa es el coeficiente de amplificación de suelo, Tc el periodo límite y r el exponente que depende del tipo de suelo.",
    "donde Fa es el coeficiente de amplificación de suelo, T0 y Tc los periodos límite de la meseta y r el exponente "
    "que depende del tipo de suelo; la rama ascendente (T < T0) se aplica a los modos distintos del fundamental.")
# Ec. (5): se normaliza por la masa total en la dirección r
rep(find("Para cada modo se calcula la razón de masa participativa"),
    "donde {r} es el vector de influencia de la dirección considerada.",
    "donde {r} es el vector de influencia de la dirección considerada y el denominador incluye la masa total en esa dirección, {r}ᵀ[M]{r}.")

d.save(DST)
print("\n".join(log)); print("guardado", DST)
